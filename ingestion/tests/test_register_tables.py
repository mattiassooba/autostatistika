"""Register general tables (üldtabelid, 2016-2018): stock, age, makes, counties."""
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook

from ingestion.discovery.transpordiamet import _classify
from ingestion.parsers.fleet import parse_register_tables


def build(path: Path) -> Path:
    wb = Workbook()
    age = wb.active
    age.title = "Yldtabel-vanus"
    age.append(["Liiklusregistris arvel olevad sõidukid. NB! Sisaldab peatatud registrikandega sõidukeid."])
    age.append(["Kuupäev: 31.12.2017"])
    age.append([None])
    age.append(["Liik", "Kuni 2a", "3 kuni 5a", "6 kuni 10a", "Üle 10a", "KOKKU"])
    age.append(["Auto Haagis", 10, 20, 30, 40, 100])
    age.append(["SELLEST ERA", 1, 2, 3, 4, 10])
    age.append(["Moto", 5, 5, 5, 5, 20])
    age.append(["Sõiduauto", 100, 200, 300, 400, 1000])
    age.append(["Traktori Haagis", 1, 1, 1, 2, 5])
    age.append(["KOKKU:", 116, 226, 336, 447, 1125])
    cars = wb.create_sheet("Sõiduautod")
    cars.append(["Liiklusregistris arvel olevad sõiduautod"])
    cars.append(["Seisuga: 31.12.2017"])
    cars.append([None])
    cars.append(["MARK", 2017, 2016, "Vanemad", "KOKKU"])
    cars.append(["TOYOTA", 10, 20, 570, 600])
    cars.append(["ŠKODA", 5, 5, 390, 400])
    cars.append(["KOKKU", 15, 25, 960, 1000])
    counties = wb.create_sheet("Maakonnad")
    counties.append(["Liiklusregistris arvel olevad sõidukid"])
    counties.append(["Seisuga 31.12.2017"])
    counties.append([None])
    counties.append([None, "Sõiduauto", None, "Mootorratas", None])
    counties.append(["Maakond", "Kokku", "sh. era", "Kokku", "sh. era"])
    counties.append(["Harju Maakond", 700, 500, 15, 14])
    counties.append(["Määramata", 300, 280, 5, 5])
    counties.append(["Kokku", 1000, 780, 20, 19])
    wb.save(path)
    return path


def test_general_tables(tmp_path: Path) -> None:
    facts = parse_register_tables(build(tmp_path / "yldtabelid-2017.xlsx"), "2017-12-31")
    by = defaultdict(float)
    for f in facts:
        by[(f.extra["grain"], f.category_code)] += f.value
    # 31.12.2017 is the register moment of 1.1.2018, like every month-end stock.
    assert {f.snapshot_date for f in facts} == {"2018-01-01"}
    assert by[("stock_total", "passenger_car")] == 1000
    assert by[("stock_total", "trailer")] == 100
    assert by[("stock_total", "tractor_trailer")] == 5
    assert by[("age", "passenger_car")] == 1000
    assert by[("brand", "passenger_car")] == 1000  # a complete split by make
    assert by[("region", "passenger_car")] == 1000
    stock = [f for f in facts if f.extra["grain"] == "stock_total"]
    assert all(f.extra["status_scope"] == "includes_suspended" for f in stock)
    # "Moto" counts motorcycles and mopeds together: not split, not published.
    assert not any(f.category_code in {"motorcycle", "moped"} for f in facts)


def test_general_tables_are_their_own_kind() -> None:
    href = "/sites/default/files/documents/2021-11/yldtabelid-2017.xlsx"
    assert _classify(href, "Arvel olevad sõidukid seisuga 31.12.2017 (Üldtabelid)")[1] == "register_tables"
    assert _classify("/x/arvel-m1_311217.xlsx", "Arvel olevad sõiduautod seisuga 31.12.2017")[1] == "fleet_snapshot"
